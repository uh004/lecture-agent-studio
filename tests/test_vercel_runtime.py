import asyncio
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi import BackgroundTasks, UploadFile
from PIL import Image
from pptx import Presentation

import main as api_main
from src.core.utils import render_slide_fallback
from src.infrastructure.job_store import JobStore
from src.infrastructure.object_storage import ObjectStorage
from src.nodes.parser import node_parse_ppt
from src.service import process_lecture_job


class VercelRuntimeAdapterTests(unittest.TestCase):
    def test_vercel_generate_publishes_source_and_enqueues_job(self) -> None:
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[1])
        content = BytesIO()
        presentation.save(content)

        with patch.dict(os.environ, {}, clear=True):
            store = JobStore()

        class RemoteStorage:
            is_remote = True

            @staticmethod
            def publish_source(data, job_id, filename):
                self.assertTrue(data)
                self.assertEqual(filename, "lecture.pptx")
                return {
                    "url": f"https://blob.example/{job_id}/lecture.pptx",
                    "download_url": "",
                    "storage": "vercel_blob",
                }

        upload = UploadFile(filename="lecture.pptx", file=BytesIO(content.getvalue()))
        with (
            patch.object(api_main, "is_vercel", True),
            patch.object(api_main, "job_store", store),
            patch.object(api_main, "object_storage", RemoteStorage()),
            patch.object(api_main, "_require_vercel_integrations"),
            patch.object(api_main, "wake_worker") as wake,
        ):
            result = asyncio.run(
                api_main.generate_video(BackgroundTasks(), upload)
            )

        job = store.get(result["job_id"])
        self.assertEqual(job["status"], "pending")
        self.assertTrue(job["source_url"].startswith("https://blob.example/"))
        self.assertEqual(store.queue_depth(), 1)
        wake.assert_called_once_with()

    def test_memory_job_store_round_trip(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            store = JobStore()
            self.assertFalse(store.is_durable)
            created = store.create("job-1", {"status": "pending"})
            self.assertEqual(created["status"], "pending")

            updated = store.update("job-1", status="running", current_slide=1)
            self.assertEqual(updated["status"], "running")
            self.assertEqual(store.get("job-1")["current_slide"], 1)

    def test_memory_job_queue_is_fifo(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            store = JobStore()
            store.enqueue("job-1")
            store.enqueue("job-2")
            self.assertEqual(store.queue_depth(), 2)
            self.assertEqual(store.dequeue(), "job-1")
            self.assertEqual(store.dequeue(), "job-2")
            self.assertIsNone(store.dequeue())

    def test_memory_job_queue_recovers_unacknowledged_job(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            store = JobStore()
            store.enqueue("job-1")
            self.assertEqual(store.dequeue(), "job-1")
            self.assertEqual(store.inflight_depth(), 1)
            self.assertEqual(store.recover_inflight(), 1)
            self.assertEqual(store.dequeue(), "job-1")
            store.acknowledge("job-1")
            self.assertEqual(store.inflight_depth(), 0)

    def test_process_job_uses_shared_state_contract(self) -> None:
        with patch.dict(os.environ, {}, clear=True), tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "lecture.pptx"
            source.write_bytes(b"pptx")
            store = JobStore()
            storage = ObjectStorage()
            store.create(
                "job-1",
                {
                    "status": "pending",
                    "source_path": str(source),
                    "original_name": source.name,
                    "settings": {"speed": 1.0},
                },
            )

            def fake_run(*args, **kwargs):
                final_video = Path(kwargs["work_dir"]) / "final_lecture.mp4"
                final_video.parent.mkdir(parents=True, exist_ok=True)
                final_video.write_bytes(b"video")
                return {
                    "final_video": str(final_video),
                    "final_status": "completed",
                    "final_qa": {"passed": True},
                    "errors": [],
                    "total_slides": 1,
                }

            with patch("src.service.run_lecture_agent", side_effect=fake_run):
                process_lecture_job(
                    "job-1", store, storage, work_root=root / "work", cleanup=False
                )

            result = store.get("job-1")
            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["storage"], "local")
            self.assertTrue(Path(result["final_video_path"]).exists())

    def test_local_object_storage_keeps_file(self) -> None:
        with patch.dict(os.environ, {}, clear=True):
            storage = ObjectStorage()
            with tempfile.TemporaryDirectory() as directory:
                video = Path(directory) / "video.mp4"
                video.write_bytes(b"not-empty")
                result = storage.publish_video(video, "job-1")
                self.assertEqual(result["storage"], "local")
                self.assertEqual(Path(result["url"]), video.resolve())

    def test_fallback_slide_renderer_creates_full_hd_png(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "slide.png"
            rendered = render_slide_fallback(
                title="모델 모니터링",
                body_texts=["데이터 변화와 성능을 지속적으로 확인합니다."],
                tables=[[['항목', '설명'], ['성능', '정확도 확인']]],
                images=[],
                output_path=str(output),
            )
            self.assertEqual(Path(rendered), output)
            self.assertTrue(output.exists())
            with Image.open(output) as image:
                self.assertEqual(image.size, (1920, 1080))

    def test_parser_falls_back_when_office_renderer_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pptx_path = root / "sample.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Vercel 슬라이드"
            slide.placeholders[1].text = "시스템 렌더러가 없어도 이미지가 생성됩니다."
            presentation.save(pptx_path)

            state = {
                "pptx_path": str(pptx_path),
                "work_dir": str(root / "work"),
                "lecture_config": {},
            }
            with patch(
                "src.nodes.parser.export_slide_as_png",
                side_effect=RuntimeError("soffice unavailable"),
            ):
                result = node_parse_ppt(state)

            rendered = Path(result["slides"][0]["slide_image"])
            self.assertTrue(rendered.exists())
            self.assertEqual(result["slides"][0]["extraction_errors"], [])


if __name__ == "__main__":
    unittest.main()
