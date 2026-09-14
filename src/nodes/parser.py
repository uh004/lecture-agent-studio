"""Node 1: PPTX를 슬라이드 단위 구조화 데이터로 파싱합니다."""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from src.core.state import AgentState
from src.core.utils import (
    clean_text,
    export_slide_as_png,
    record_error,
    render_slide_fallback,
    reset_slide_runtime,
    safe_print,
)


print = safe_print


SLIDE_REQUIRED_KEYS = {
    "index", "title", "text", "notes", "tables", "charts",
    "images", "slide_image", "shape_texts", "links", "extraction_errors",
}


def node_parse_ppt(state: AgentState) -> AgentState:
    print("\n--- Node 1: PPT 파싱 실행 ---")
    if not state.get("lecture_config") and isinstance(state.get("prompt"), dict):
        state["lecture_config"] = dict(state["prompt"])
    pptx_path = Path(state["pptx_path"]).expanduser().resolve()
    work_dir = Path(state.get("work_dir", "./output_v4")).expanduser().resolve()
    slides_dir = work_dir / "slides"
    media_dir = work_dir / "media"
    slides_dir.mkdir(parents=True, exist_ok=True)
    media_dir.mkdir(parents=True, exist_ok=True)

    state.setdefault("errors", [])
    state.setdefault("failed_slides", [])
    state.setdefault("failed_slide_details", [])
    state.setdefault("all_scripts", [])
    state.setdefault("video_paths", [])
    state.setdefault("max_validation_attempts", 2)
    state.setdefault("max_validation_error_retries", 1)
    state["final_video"] = ""
    state["final_qa"] = {}
    state["final_status"] = "running"

    if not pptx_path.exists():
        message = f"PPTX 파일이 존재하지 않습니다: {pptx_path}"
        record_error(state, message)
        state["final_status"] = "failed"
        raise FileNotFoundError(message)

    try:
        presentation = Presentation(str(pptx_path))
    except Exception as exc:
        message = f"PPTX 열기 실패: {exc}"
        record_error(state, message)
        state["final_status"] = "failed"
        raise RuntimeError(message) from exc

    slides: List[Dict[str, Any]] = []
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    url_pattern = r"https?://[^\s]+"

    for slide_index, slide in enumerate(presentation.slides):
        print(f"  슬라이드 {slide_index + 1}/{len(presentation.slides)} 추출")
        slide_errors: List[str] = []
        slide_image = ""
        native_render_error = ""
        destination = slides_dir / f"slide_{slide_index + 1}.png"
        try:
            rendered = export_slide_as_png(str(pptx_path), str(slides_dir), slide_index)
            if Path(rendered).resolve() != destination.resolve():
                os.replace(rendered, destination)
            slide_image = str(destination)
        except Exception as exc:
            native_render_error = str(exc)

        notes = ""
        try:
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame:
                    notes = clean_text(notes_frame.text)
        except Exception as exc:
            message = f"슬라이드 {slide_index + 1} 발표자 노트 추출 실패: {exc}"
            slide_errors.append(message)
            record_error(state, message)

        title = ""
        for shape in slide.shapes:
            try:
                if shape.is_placeholder and shape.placeholder_format.type in title_types:
                    if shape.has_text_frame and clean_text(shape.text):
                        title = clean_text(shape.text)
                        break
            except Exception:
                continue

        body_texts: List[str] = []
        shape_texts: List[str] = []
        links = set()
        tables: List[List[List[str]]] = []
        charts: List[Dict[str, Any]] = []
        images: List[str] = []

        def collect_shape_text(shape: Any) -> None:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    collect_shape_text(child)
                return
            if getattr(shape, "has_text_frame", False):
                value = clean_text(shape.text_frame.text)
                if value and value != title and value not in shape_texts:
                    shape_texts.append(value)

        def collect_images(shape: Any) -> None:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    collect_images(child)
                return
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                return
            try:
                extension = shape.image.ext or "png"
                image_path = media_dir / f"slide{slide_index + 1}_img_{len(images) + 1}.{extension}"
                image_path.write_bytes(shape.image.blob)
                images.append(str(image_path))
            except Exception as exc:
                message = f"슬라이드 {slide_index + 1} 이미지 추출 실패: {exc}"
                slide_errors.append(message)
                record_error(state, message)

        for shape in slide.shapes:
            collect_shape_text(shape)
            collect_images(shape)

            if getattr(shape, "has_table", False):
                try:
                    table = [[clean_text(cell.text) for cell in row.cells] for row in shape.table.rows]
                    tables.append(table)
                except Exception as exc:
                    message = f"슬라이드 {slide_index + 1} 표 추출 실패: {exc}"
                    slide_errors.append(message)
                    record_error(state, message)

            if getattr(shape, "has_chart", False):
                try:
                    chart = shape.chart
                    chart_title = ""
                    if chart.has_title and chart.chart_title.has_text_frame:
                        chart_title = clean_text(chart.chart_title.text_frame.text)
                    series_names = [clean_text(series.name) for series in chart.series]
                    categories = []
                    if chart.plots:
                        categories = [clean_text(category.label) for category in chart.plots[0].categories]
                    rows = []
                    for category_index, category in enumerate(categories):
                        values = []
                        for series in chart.series:
                            values.append(
                                series.values[category_index]
                                if category_index < len(series.values)
                                else None
                            )
                        rows.append({"category": category, "values": values})
                    charts.append({
                        "title": chart_title,
                        "series": series_names,
                        "categories": categories,
                        "rows": rows,
                    })
                except Exception as exc:
                    message = f"슬라이드 {slide_index + 1} 차트 추출 실패: {exc}"
                    slide_errors.append(message)
                    record_error(state, message)

            if getattr(shape, "has_text_frame", False):
                try:
                    is_title = shape.is_placeholder and shape.placeholder_format.type in title_types
                except Exception:
                    is_title = False
                if not is_title:
                    for paragraph in shape.text_frame.paragraphs:
                        paragraph_text = clean_text("".join(run.text for run in paragraph.runs))
                        if paragraph_text and paragraph_text not in body_texts:
                            body_texts.append(paragraph_text)
                        for run in paragraph.runs:
                            try:
                                address = run.hyperlink.address
                                if address:
                                    links.add(address)
                            except Exception:
                                pass
                        links.update(re.findall(url_pattern, paragraph_text))

            try:
                address = shape.click_action.hyperlink.address
                if address:
                    links.add(address)
            except Exception:
                pass

        if not slide_image:
            try:
                slide_image = render_slide_fallback(
                    title=title,
                    body_texts=body_texts,
                    tables=tables,
                    images=images,
                    output_path=str(destination),
                )
                print(
                    f"  ⚠️ 슬라이드 {slide_index + 1}: 시스템 렌더러 대신 "
                    f"Vercel 호환 렌더러 사용 ({native_render_error})"
                )
            except Exception as exc:
                message = f"슬라이드 {slide_index + 1} PNG 생성 실패: {exc}"
                slide_errors.append(message)
                record_error(state, message)

        slides.append({
            "index": slide_index,
            "title": title,
            "text": "\n".join(body_texts).strip(),
            "notes": notes,
            "tables": tables,
            "charts": charts,
            "images": images,
            "slide_image": slide_image,
            "shape_texts": shape_texts,
            "links": sorted(links),
            "extraction_errors": slide_errors,
        })

    state["slides"] = slides
    state["total_slides"] = len(slides)
    state["slide_index"] = 0
    reset_slide_runtime(state)
    print(f"✅ PPT 파싱 완료: {len(slides)}개 슬라이드")
    return state


def validate_slide_contract(slide: Dict[str, Any], expected_index: int) -> None:
    missing = SLIDE_REQUIRED_KEYS - set(slide)
    assert not missing, f"슬라이드 필수 키 누락: {sorted(missing)}"
    assert slide["index"] == expected_index, (
        f"슬라이드 인덱스 불일치: {slide['index']} != {expected_index}"
    )


def validate_parser_state(state: AgentState) -> None:
    slides = state.get("slides", [])
    assert state.get("total_slides") == len(slides), "전체 슬라이드 수 불일치"
    assert slides, "추출된 슬라이드가 없습니다."
    for expected_index, slide in enumerate(slides):
        validate_slide_contract(slide, expected_index)
